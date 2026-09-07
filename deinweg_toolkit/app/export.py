"""Export der gefilterten Zeiten als xlsx und csv.

Eigener Router, damit main.py nicht auch noch openpyxl und csv tragen
muss. Der Filter selbst kommt aus rechnen.py - dieselbe Funktion, die
Auswertung und Uebersicht benutzen; sonst zeigte der Export etwas
anderes an als die Liste, aus der man ihn anstoesst.
"""

from __future__ import annotations

import csv
import io
import unicodedata
from urllib.parse import quote

from fastapi import APIRouter, Query
from fastapi.responses import StreamingResponse

from . import db
from .parser import hhmm
from .rechnen import MONATSNAMEN, bereichsfilter, deutsch, sicherer_name

router = APIRouter()


SPALTEN = ["Datum", "Betreuter", "Beginn", "Ende", "Dauer",
           "Leistung/Beschreibung", "Mitarbeiter"]


def hole_zeilen(wo, werte):
    with db.db() as con:
        return con.execute(
            f"SELECT * FROM eintrag WHERE {wo} ORDER BY datum, mitarbeiter, start",
            werte).fetchall()


def baue_xlsx(zeilen) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Nachweis"
    ws.append(SPALTEN)
    for zelle in ws[1]:
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = PatternFill("solid", fgColor="272827")
        zelle.alignment = Alignment(vertical="center")
    for z in zeilen:
        ws.append([deutsch(z["datum"]), z["klient"], z["start"], z["ende"],
                   hhmm(z["dauer_min"]), z["beschreibung"], z["mitarbeiter"]])
    summe = sum(z["dauer_min"] for z in zeilen)
    ws.append([])
    ws.append(["Gesamt", "", "", "", hhmm(summe), f"{len(zeilen)} Einträge", ""])
    for zelle in ws[ws.max_row]:
        zelle.font = Font(bold=True)
    for spalte, breite in zip("ABCDEFG", [12, 24, 9, 9, 10, 54, 16]):
        ws.column_dimensions[spalte].width = breite
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(2, len(zeilen) + 1)}"

    puffer = io.BytesIO()
    wb.save(puffer)
    puffer.seek(0)
    return puffer


def exportname(filter_: dict, endung: str) -> str:
    """Dateiname aus dem gewählten Zeitraum und Mitarbeiter."""
    von, bis = filter_["von"], filter_["bis"]
    nur_a, nur_b = filter_.get("nur_monate", ("", ""))
    if von and bis and von == bis:
        zeit = von
    elif von and bis:
        zeit = f"{von}_bis_{bis}"
    elif von:
        zeit = f"ab_{von}"
    elif bis:
        zeit = f"bis_{bis}"
    elif nur_a or nur_b:
        # Monat ohne Jahreszahl: Monatsname statt Datumsspanne
        namen = [MONATSNAMEN.get(x, x) for x in (nur_a, nur_b) if x]
        zeit = ("-".join(dict.fromkeys(namen))) + "_alle-Jahre"
    else:
        zeit = "gesamt"
    teile = ["Zeitnachweis", zeit]
    if filter_["f"]["mitarbeiter"]:
        teile.append(sicherer_name(filter_["f"]["mitarbeiter"]).replace(" ", "-"))
    return "_".join(teile) + endung


def dateikopf(name: str) -> str:
    """Der Content-Disposition-Kopf zu einem Dateinamen.

    ⚠️ HTTP-Koepfe tragen keinen Unicode. Bis 1.32 stand der Name roh im
    Kopf - bei „Zeitnachweis_März-Juni_alle-Jahre.csv" oder einem
    Mitarbeiter namens „Müller" kam beim Browser Buchstabensalat an, und
    bei einem Zeichen ausserhalb von Latin-1 waere die Antwort mit einem
    Serverfehler abgebrochen. Aufgefallen ist es erst, als der Export
    beim Aufteilen von main.py eine eigene Pruefung bekam.

    Deshalb zweimal: ein ASCII-Name fuer alte Programme und derselbe
    Name als filename* nach RFC 6266, den jeder heutige Browser
    bevorzugt.
    """
    roh = unicodedata.normalize("NFKD", name)
    ascii_name = roh.encode("ascii", "ignore").decode("ascii") or "export"
    return (f'attachment; filename="{ascii_name}"; '
            f"filename*=UTF-8''{quote(name)}")


@router.get("/export.xlsx")
def export_xlsx(von_jahr: str = "", von_monat: str = "", bis_jahr: str = "",
                bis_monat: str = "", mitarbeiter: list[str] = Query([]),
                klient: list[str] = Query([]),
                q: str = "", import_id: int = 0, nur_abrechenbar: str = ""):
    filter_ = bereichsfilter(von_jahr, von_monat, bis_jahr, bis_monat,
                             mitarbeiter, klient, q, import_id, nur_abrechenbar)
    puffer = baue_xlsx(hole_zeilen(filter_["wo"], filter_["werte"]))
    name = exportname(filter_, ".xlsx")
    return StreamingResponse(
        puffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": dateikopf(name)})


@router.get("/export.csv")
def export_csv(von_jahr: str = "", von_monat: str = "", bis_jahr: str = "",
               bis_monat: str = "", mitarbeiter: list[str] = Query([]),
               klient: list[str] = Query([]),
               q: str = "", import_id: int = 0, nur_abrechenbar: str = ""):
    filter_ = bereichsfilter(von_jahr, von_monat, bis_jahr, bis_monat,
                             mitarbeiter, klient, q, import_id, nur_abrechenbar)
    puffer = io.StringIO()
    schreiber = csv.writer(puffer, delimiter=";")
    schreiber.writerow(SPALTEN)
    for z in hole_zeilen(filter_["wo"], filter_["werte"]):
        schreiber.writerow([deutsch(z["datum"]), z["klient"], z["start"], z["ende"],
                            hhmm(z["dauer_min"]), z["beschreibung"],
                            z["mitarbeiter"]])
    daten = ("\ufeff" + puffer.getvalue()).encode("utf-8")
    name = exportname(filter_, ".csv")
    return StreamingResponse(io.BytesIO(daten), media_type="text/csv",
                             headers={"Content-Disposition": dateikopf(name)})
